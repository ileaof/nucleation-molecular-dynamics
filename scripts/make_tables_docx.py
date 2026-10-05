# -*- coding: utf-8 -*-
"""
Side-by-side tables (continuum × molecular dynamics) for Al-0.8Si-0.6Mg-0.2Fe, in the format
of the manuscript (Full_Manuscript_Phase_Nucleation_edited_R2.docx, Tables 1-2) and of the
properties table "Tabela 1 - Propriedades simulação Luane Final.docx".

    python scripts/make_tables_docx.py  ->  results/tables/Tables_AlSiMgFe_continuum_vs_MD.docx

Continuum values: data/continuum/Al08Si06Mg02Fe/tablefull_AlSiMgFe.csv (author's nucleation
script, scripts/tablefull_alsimgfe.py). MD values: results/md/properties_md.json (null = running).
"""
import json
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parents[1]
TABLE = ROOT / "data" / "continuum" / "Al08Si06Mg02Fe" / "tablefull_AlSiMgFe.csv"
MDJSON = ROOT / "results" / "md" / "properties_md.json"
OUT = ROOT / "results" / "tables" / "Tables_AlSiMgFe_continuum_vs_MD.docx"
ALLOY = "Al-0.8wt%Si-0.6wt%Mg-0.2wt%Fe"
GRAD_EXP = 7902.38
ROWS = [1, 16, 33, 44, 53, 59, 65, 70, "exp", 74, 77]
PENDING = "em cálculo"

T1 = [("dTdr", "∇T", "K·m⁻¹"), ("DT", "ΔT", "K"), ("r_hom_1st", "r_C,Hom 1st", "m"),
      ("r_het_1st", "r_C,Het 1st", "m"), ("theta_1st", "θ 1st", "rad"), ("dfthetadr_1st", "∂f(θ 1st)/∂r", "1/m"),
      ("r_hom_2nd", "r_C,Hom 2nd", "m"), ("r_het_2nd", "r_C,Het 2nd", "m"), ("theta_2nd", "θ 2nd", "rad"),
      ("dfthetadr_2nd", "∂f(θ 2nd)/∂r", "1/m"), ("sigma", "σ_SL", "N·m⁻¹"), ("surface_stress", "Σ", "N·m⁻¹"),
      ("gam_hom", "γ_SL", "J·m⁻²")]
T2 = [("dTdr", "∇T", "K·m⁻¹"), ("DT", "ΔT", "K"), ("dfgamdr_ana", "∂γ_SL/∂r", "J·m⁻³"),
      ("DSv_hom", "ΔS_V", "J·m⁻³K⁻¹"), ("GT_het", "Γ 1st", "m·K"), ("GT_het_2nd", "Γ 2nd", "m·K"),
      ("DSv_het", "ΔS_V", "J·m⁻³K⁻¹"), ("dDSv_homdr", "∂ΔS_V/∂r", "J·m⁻⁴K⁻¹"), ("DSs_hom", "ΔS_S", "J·m⁻³K⁻¹"),
      ("DSc", "ΔS_C", "J·m⁻³K⁻¹"), ("GB", "ΔG_V", "J·m⁻³"), ("GS", "ΔG_S", "J·m⁻²"), ("GC", "ΔG_C", "J·m⁻³")]
SCI = {"r_hom_1st", "r_het_1st", "r_hom_2nd", "r_het_2nd", "GT_het", "GT_het_2nd", "dDSv_homdr"}


def fmt(col, x):
    if isinstance(x, str):
        return x
    return f"{x:.5E}" if col in SCI else f"{x:.6g}"


def table_rows():
    t = pd.read_csv(TABLE, sep=";")
    out = []
    for r in ROWS:
        if r == "exp":
            # linear interpolation in dTdr between the bracketing rows (as the script, L1036)
            j = t.index[t.dTdr <= GRAD_EXP].max()
            a, b = t.loc[j], t.loc[j + 1]
            w = (GRAD_EXP - a.dTdr) / (b.dTdr - a.dTdr)
            row = a + w * (b - a)
            row["dTdr"] = GRAD_EXP
            out.append((row, True))
        else:
            out.append((t[t.i == r].iloc[0], False))
    return out


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def put(cell, text, bold=False, size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER, italic=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    run.bold, run.italic = bold, italic
    run.font.size = Pt(size)


def caption(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.size = Pt(10)


def nucleation_table(doc, title, cols, rows):
    caption(doc, title)
    tb = doc.add_table(rows=2 + len(rows) + 2, cols=len(cols))
    tb.style = "Table Grid"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, (_, sym, unit) in enumerate(cols):
        put(tb.cell(0, k), sym, bold=True); put(tb.cell(1, k), unit, italic=True)
        shade(tb.cell(0, k), "E8E8E8"); shade(tb.cell(1, k), "E8E8E8")
    for r, (row, is_exp) in enumerate(rows):
        for k, (c, _, _) in enumerate(cols):
            txt = fmt(c, row[c]) + ("*" if is_exp and k == 0 else "")
            put(tb.cell(2 + r, k), txt, bold=is_exp)
    # MD block (atomistic scale) — filled when seeding/NEMD results exist
    r0 = 2 + len(rows)
    m = tb.cell(r0, 0).merge(tb.cell(r0, len(cols) - 1))
    put(m, "Dinâmica molecular (escala atomística, r_C ~ nm, ∇T ~ 10⁸–10⁹ K·m⁻¹) — Borovikov et al. (2024) EAM/FS",
        bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    shade(m, "F2F2F2")
    m2 = tb.cell(r0 + 1, 0).merge(tb.cell(r0 + 1, len(cols) - 1))
    put(m2, f"{PENDING}: seeding + NEMD (etapa na GPU)", italic=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    p = doc.add_paragraph()
    r = p.add_run(f"*Gradiente térmico médio experimental (Marques et al.), linha interpolada linearmente em ∇T. "
                  f"Contínuo: script de nucleação do autor ({ALLOY}).")
    r.font.size = Pt(8); r.italic = True


def md(js, key, fmtspec="{:g}"):
    v = js.get(key, {}).get("value")
    if v is None:
        return PENDING
    unc = js[key].get("unc")
    return fmtspec.format(v) + (f" ± {unc:g}" if unc else "")


def properties_table(doc, js, cont):
    caption(doc, f"Tabela 3 – Propriedades termofísicas da liga {ALLOY}: mecânica do contínuo × dinâmica molecular")
    exp = cont
    rows = [
        ("Temperatura liquidus / de fusão", "T_L", "K", "925.32", md(js, "T_m_K") + " (Al)"),
        ("Temperatura solidus", "T_SOL", "K", "798.15", "—"),
        ("Inclinação liquidus (Mg, Si, Fe)", "m_L", "K (wt%)⁻¹", "−5.100; −6.200; −4.200", "—"),
        ("Coeficiente de partição (Mg, Si, Fe)", "k₀", "–", "0.36; 0.11; 0.03", "—"),
        ("Calor latente (FCC_A1)", "ΔH_FCC_A1", "J·kg⁻¹", "335300", md(js, "dH_m_J_per_kg", "{:.0f}")),
        ("Massa específica do sólido em T_SOL", "ρ_s", "kg·m⁻³", "2555.72", md(js, "rho_s_Tsol", "{:.1f}")),
        ("Massa específica do líquido em T_SOL", "ρ_l", "kg·m⁻³", "2378.33", md(js, "rho_l_Tsol", "{:.1f}")),
        ("Gradiente térmico médio (líquido)", "∇T", "K·m⁻¹", "7902.38", "—"),
        ("Tensão superficial a baixo gradiente", "σ₀", "N·m⁻¹", "1.09", "—"),
        ("Energia de superfície a baixo gradiente", "γ₀", "J·m⁻²", "0.183", md(js, "gamma0_J_per_m2", "{:.4f}")),
        ("Tensão de superfície em ∇T exp.", "Σ", "N·m⁻¹", f"{exp['surface_stress']:.4g}", "—"),
        ("Tensão superficial em ∇T exp.", "σ_SL", "N·m⁻¹", f"{exp['sigma']:.4g}", "—"),
        ("Energia de superfície em ∇T exp.", "γ_SL", "J·m⁻²", f"{exp['gam_hom']:.4g}", "—"),
        ("Gibbs–Thomson 1ª ordem", "Γ_Het 1st", "m·K", f"{exp['GT_het']:.4E}", "—"),
        ("Gibbs–Thomson 2ª ordem", "Γ_Het 2nd", "m·K", f"{exp['GT_het_2nd']:.4E}", md(js, "Gamma_2nd_mK", "{:.4E}")),
        ("Entropia volumétrica", "ΔS_V = −ΔH·ρ_s/T", "J·m⁻³K⁻¹", f"{-335300*2555.72/925.32:.4E}", "−8.878E+05 (T_m, Al)"),
    ]
    tb = doc.add_table(rows=1 + len(rows), cols=5)
    tb.style = "Table Grid"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, h in enumerate(("Propriedade", "Símbolo", "Unidade", "Mecânica do contínuo", "Dinâmica molecular")):
        put(tb.cell(0, k), h, bold=True, size=9); shade(tb.cell(0, k), "E8E8E8")
    for r, row in enumerate(rows, start=1):
        for k, v in enumerate(row):
            put(tb.cell(r, k), v, size=9, align=WD_ALIGN_PARAGRAPH.LEFT if k == 0 else WD_ALIGN_PARAGRAPH.CENTER)
    p = doc.add_paragraph()
    r = p.add_run("Contínuo: scripts de nucleação e de SDAS do autor (σ₀ = 1,09 N·m⁻¹ e γ₀ = 0,183 J·m⁻² são os valores "
                  "que geram Γ). MD: Al com o potencial EAM/FS de Borovikov et al. (2024); o potencial não contém Si, Mg "
                  "nem Fe, por isso as propriedades de soluto ficam “—”. ρ em T_SOL na MD: sólido e líquido "
                  "super-resfriado em NPT.")
    r.font.size = Pt(8); r.italic = True


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    js = json.loads(MDJSON.read_text(encoding="utf-8"))
    rows = table_rows()
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.5))
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(10)
    nucleation_table(doc, f"Tabela 1 – Variáveis de nucleação calculadas em função do gradiente térmico para a fase "
                          f"FCC_A1. Composição da liga: {ALLOY}.", T1, rows)
    doc.add_page_break()
    nucleation_table(doc, f"Tabela 2 – Variáveis de nucleação calculadas em função do gradiente térmico para a fase "
                          f"FCC_A1. Composição da liga: {ALLOY}.", T2, rows)
    doc.add_page_break()
    properties_table(doc, js, [r for r, e in rows if e][0])
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
